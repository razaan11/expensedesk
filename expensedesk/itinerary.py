"""Build a clean travel itinerary for a partner/director from structured legs."""
from io import BytesIO
from docx import Document


def _sorted(legs):
    return sorted(legs, key=lambda x: (x["date"], x["time"]))


def build_itinerary_text(traveller: str, trip_title: str, legs: list[dict], hotel: dict | None = None,
                         notes: str = "") -> str:
    lines = [f"ITINERARY: {trip_title}", f"Traveller: {traveller}", ""]
    for i, leg in enumerate(_sorted(legs), 1):
        lines.append(f"{i}. {leg['date']} {leg['time']} | {leg['mode']} | {leg['from']} -> {leg['to']} | Ref: {leg.get('ref', 'TBC')}")
    if hotel:
        lines += ["", f"Hotel: {hotel['name']}, {hotel['city']} | {hotel['check_in']} to {hotel['check_out']} | Ref: {hotel.get('ref', 'TBC')}"]
    if notes:
        lines += ["", f"Notes: {notes}"]
    return "\n".join(lines)


def build_itinerary_docx(traveller: str, trip_title: str, legs: list[dict], hotel: dict | None = None,
                         notes: str = "") -> bytes:
    doc = Document()
    doc.add_heading(f"Itinerary: {trip_title}", level=1)
    doc.add_paragraph(f"Traveller: {traveller}")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, h in zip(table.rows[0].cells, ["Date", "Time", "Mode", "Route", "Reference"]):
        cell.text = h
    for leg in _sorted(legs):
        cells = table.add_row().cells
        for cell, v in zip(cells, [leg["date"], leg["time"], leg["mode"], f"{leg['from']} -> {leg['to']}", leg.get("ref", "TBC")]):
            cell.text = str(v)
    if hotel:
        p = doc.add_paragraph()
        p.add_run("Hotel: ").bold = True
        p.add_run(f"{hotel['name']}, {hotel['city']} ({hotel['check_in']} to {hotel['check_out']}) Ref: {hotel.get('ref', 'TBC')}")
    if notes:
        doc.add_paragraph(f"Notes: {notes}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
